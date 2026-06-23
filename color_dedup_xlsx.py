"""Rebuild the dedup workbook from the (media-filled) CSV with two sheets and
cluster shading:
  • 'flagged'  — full data; rows shaded by dup_cluster_id parity (even=light
    blue, odd=light orange; non-cluster rows unfilled), bold + frozen header.
  • 'clusters' — one row per duplicate cluster (size, years, canonical title).
"""
import re
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font

CSV = "PURSUE_1_2_3_dedup_flagged.csv"
XLSX = "PURSUE_1_2_3_dedup_flagged.xlsx"
EVEN = PatternFill("solid", fgColor="DCE6F1")   # light blue
ODD = PatternFill("solid", fgColor="FDE9D9")    # light orange
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _san(df):
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].map(lambda v: _ILLEGAL.sub("", v) if isinstance(v, str) else v)
    return df


df = _san(pd.read_csv(CSV, low_memory=False))
# sort by decreasing dup_cluster_id (clusters grouped, descending; singletons last)
df = df.sort_values("dup_cluster_id", ascending=False, na_position="last",
                    kind="stable").reset_index(drop=True)

# reconstruct the clusters summary from the data
cl = df[df["dup_cluster_id"].notna()].copy()
rows = []
for cid, g in cl.groupby("dup_cluster_id"):
    canon = g[g["dup_role"] == "canonical"]
    title = ""
    if len(canon):
        title = canon["Title"].iloc[0]
        if pd.isna(title):
            title = canon["document_id"].iloc[0]
    rows.append({
        "cluster_id": int(cid),
        "size": len(g),
        "year(s)": ",".join(sorted({str(int(y)) for y in g["date_time.year"].dropna()})),
        "roles": f"{int((g.dup_role=='canonical').sum())} canonical / {int((g.dup_role=='duplicate').sum())} dup",
        "canonical_title": str(title)[:70],
    })
clusters_df = pd.DataFrame(rows).sort_values("cluster_id", ascending=False)

with pd.ExcelWriter(XLSX, engine="openpyxl") as xw:
    df.to_excel(xw, sheet_name="flagged", index=False)
    clusters_df.to_excel(xw, sheet_name="clusters", index=False)

# shade by parity
wb = openpyxl.load_workbook(XLSX)
ws = wb["flagged"]
col = {c.value: c.column for c in ws[1]}["dup_cluster_id"]
ncol = ws.max_column
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "A2"
colored = 0
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, col).value
    if v in (None, ""):
        continue
    try:
        cid = int(v)
    except (TypeError, ValueError):
        continue
    fill = EVEN if cid % 2 == 0 else ODD
    for cc in range(1, ncol + 1):
        ws.cell(r, cc).fill = fill
    colored += 1
wb.save(XLSX)
print(f"{XLSX}: sheets={wb.sheetnames} | flagged {df.shape} | clusters {clusters_df.shape[0]} | shaded {colored} rows")
