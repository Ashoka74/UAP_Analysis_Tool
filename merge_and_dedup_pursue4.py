"""Merge PURSUE_4_normalized into the PURSUE_1_2_3 corpus, then apply the
same tiered dedup pipeline as apply_dedup.py + color_dedup_xlsx.py (embed
cosine blocking + gates -> Gemini Tier-3 confirmation -> union-find
clustering -> shaded flagged workbook), extended to catch PURSUE_4
documents that re-describe an event already recorded in an earlier release.

PURSUE_4_normalized.csv lacks location_country_iso/state_norm,
craft_primary_shape_norm/size_band, witness_primary_role — the columns
geo_dedup._field_agreement relies on. For any candidate pair touching a
PURSUE_4 row, use_field_score is dropped (narrative-embedding cosine +
date/location gates carry the decision instead); PURSUE_1_2_3-only pairs
keep the full field-agreement gate, unchanged from apply_dedup.py.

Nothing is deleted; source files are untouched. Outputs:
    SUBDATASETS_V2/PURSUE_1_2_3_4_merged.csv   -- concatenated source
    narr_emb_pursue_1_2_3_4.npy                -- embeddings for the merged file
    PURSUE_1_2_3_4_confirmed_dup_pairs.csv     -- Gemini Tier-3 cache
    PURSUE_1_2_3_4_dedup_flagged.csv / .xlsx   -- flagged + shaded workbook
"""
import os
import itertools
from collections import defaultdict

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font

import make_timeline as mt
import geo_dedup as gd
import eval_pairs as ep
from apply_dedup import UF, _sanitize

SRC_123 = mt.SRC
SRC_4 = "SUBDATASETS_V2/PURSUE_4_normalized.csv"
MERGED_CSV = "SUBDATASETS_V2/PURSUE_1_2_3_4_merged.csv"
EMB_OUT = "narr_emb_pursue_1_2_3_4.npy"
OUT_PAIRS = "PURSUE_1_2_3_4_confirmed_dup_pairs.csv"
OUT_XLSX = "PURSUE_1_2_3_4_dedup_flagged.xlsx"
APPLY_EMBED_MIN = 0.80
EMBED_MODEL = "all-MiniLM-L6-v2"

EVEN = PatternFill("solid", fgColor="DCE6F1")   # light blue
ODD = PatternFill("solid", fgColor="FDE9D9")    # light orange


def build_merged():
    if os.path.exists(MERGED_CSV):
        return pd.read_csv(MERGED_CSV, low_memory=False)
    a = pd.read_csv(SRC_123, low_memory=False)
    b = pd.read_csv(SRC_4, low_memory=False)
    a["_source_dataset"] = "PURSUE_1_2_3"
    b["_source_dataset"] = "PURSUE_4"
    merged = pd.concat([a, b], ignore_index=True, sort=False)
    merged.to_csv(MERGED_CSV, index=False)
    print(f"merged {len(a)} + {len(b)} = {len(merged)} rows -> {MERGED_CSV}")
    return merged


def build_embeddings(df):
    if os.path.exists(EMB_OUT):
        emb = np.load(EMB_OUT)
        if len(emb) == len(df):
            return emb
        print("cached embeddings size mismatch — recomputing")
    from sentence_transformers import SentenceTransformer
    texts = df[mt._NARR].fillna("").astype(str).tolist()
    dev = "cpu"
    try:
        import torch
        dev = "cuda" if torch.cuda.is_available() else "cpu"
    except Exception:
        pass
    m = SentenceTransformer(EMBED_MODEL, device=dev)
    emb = m.encode(texts, batch_size=256, normalize_embeddings=True,
                   show_progress_bar=False).astype("float32")
    np.save(EMB_OUT, emb)
    print(f"embedded {len(texts)} narratives -> {EMB_OUT} (device={dev})")
    return emb


def _judge_chunked(df, cands, max_retries=3):
    """Tier-3 confirmation with per-chunk retry — eval_pairs.judge() has no
    retry/checkpointing, so one malformed Gemini response mid-run loses every
    already-judged chunk. Retries a failing chunk up to max_retries times;
    if it still fails, skips it (those pairs get no verdict) instead of
    crashing the whole batch."""
    import time
    import json as _json
    import re as _re
    import tomllib
    with open(".streamlit/secrets.toml", "rb") as f:
        key = tomllib.load(f)["GEMINI_KEY"]
    import google.generativeai as genai
    genai.configure(api_key=key)
    model = genai.GenerativeModel(ep.MODEL, generation_config={
        "response_mime_type": "application/json", "temperature": 0.0,
        "max_output_tokens": 65536})

    labels = {}
    CHUNK = 50
    for c0 in range(0, len(cands), CHUNK):
        chunk = cands[c0:c0 + CHUNK]
        blocks = [
            f"=== PAIR {k} ===\nRecord A:\n{ep._ctx(df.loc[a])}\n\nRecord B:\n{ep._ctx(df.loc[b])}"
            for k, (a, b, _s) in enumerate(chunk, start=c0)
        ]
        prompt = (
            "You are a UFO-report deduplication judge. For each numbered pair, decide "
            "whether records A and B describe the SAME real-world sighting/event — the "
            "same object seen by the same witness(es) at the same place and time — "
            "allowing for differences in wording, date/location precision, missing "
            "fields, or one source quoting another. Distinct sightings on the same day "
            "in the same region (a 'flap') are NOT the same event.\n\n"
            f"Return ONLY a JSON array of exactly {len(chunk)} objects:\n"
            '{"id": <pair id>, "same_event": <true|false>, "reason": "<short>"}\n\n'
            + "\n\n".join(blocks)
        )
        for attempt in range(1, max_retries + 1):
            try:
                resp = model.generate_content(prompt)
                try:
                    arr = _json.loads(resp.text)
                except _json.JSONDecodeError:
                    arr = _json.loads(_re.search(r"\[.*\]", resp.text, _re.DOTALL).group(0))
                for v in arr:
                    labels[int(v["id"])] = (bool(v.get("same_event", False)), str(v.get("reason", "")))
                break
            except Exception as e:
                if attempt == max_retries:
                    print(f"  ✗ chunk {c0}-{c0+len(chunk)} failed after {max_retries} "
                          f"attempts ({e}) — skipping {len(chunk)} pairs")
                else:
                    print(f"  ⚠ chunk {c0}-{c0+len(chunk)} attempt {attempt} failed ({e}) — retrying")
                    time.sleep(3 * attempt)
        print(f"  judged {min(c0+CHUNK, len(cands))}/{len(cands)}")
    return labels


def _is_pursue4(row):
    return row.get("_source_dataset") == "PURSUE_4"


def candidate_pairs(df, emb):
    dd = df[df["date_time.year"].notna() & (df["date_time.year"] > 1000)]
    by_ym = defaultdict(list)
    for i, r in dd.iterrows():
        y = mt._int_or_none(r["date_time.year"])
        m = mt._int_or_none(r.get("date_time.month"))
        by_ym[(y, m)].append(i)
    cands, n_cross = [], 0
    for v in by_ym.values():
        for a, b in itertools.combinations(v, 2):
            es = float(emb[a] @ emb[b])
            if es < APPLY_EMBED_MIN:
                continue
            ra, rb = df.loc[a], df.loc[b]
            crosses_4 = _is_pursue4(ra) or _is_pursue4(rb)
            params = dict(gd.PREDICT_PARAMS)
            if crosses_4:
                params["use_field_score"] = False   # PURSUE_4 lacks the enrichment cols
            s = mt._narrative_sim(str(ra.get(mt._NARR, "")), str(rb.get(mt._NARR, "")))
            if gd.predict_same_event(ra, rb, s, embed_sim=es, P=params):
                cands.append((a, b, round(es, 4)))
                if crosses_4:
                    n_cross += 1
    print(f"high-confidence candidate pairs (embed>={APPLY_EMBED_MIN}): {len(cands)} "
          f"(of which {n_cross} touch a PURSUE_4 row)")
    return cands


def main():
    df = build_merged()
    emb = build_embeddings(df)

    if os.path.exists(OUT_PAIRS) and os.path.getsize(OUT_PAIRS) > 50:
        cdf = pd.read_csv(OUT_PAIRS)
        confirmed = [(int(r.i_id), int(r.j_id), float(r.embed_sim), str(r.llm_reason))
                     for r in cdf.itertuples()]
        print(f"loaded {len(confirmed)} LLM-confirmed pairs from cache (skipping LLM)")
    else:
        cands = candidate_pairs(df, emb)
        labels = _judge_chunked(df, cands)           # Tier-3: LLM confirmation
        sim_by_pair = {k: cands[k][2] for k in range(len(cands))}
        confirmed = [(cands[k][0], cands[k][1], sim_by_pair[k], r)
                     for k, (same, r) in labels.items() if same]
        print(f"LLM-confirmed same-event pairs: {len(confirmed)} / {len(cands)}")
        pd.DataFrame([{"i_id": a, "j_id": b, "embed_sim": s, "llm_reason": r}
                      for a, b, s, r in confirmed]).to_csv(OUT_PAIRS, index=False)

    # cluster confirmed pairs
    uf = UF()
    sim_edge = {}
    for a, b, s, _r in confirmed:
        uf.union(a, b)
        sim_edge[a] = max(sim_edge.get(a, 0), s)
        sim_edge[b] = max(sim_edge.get(b, 0), s)
    groups = defaultdict(list)
    for a, b, _s, _r in confirmed:
        groups[uf.find(a)].append(a); groups[uf.find(b)].append(b)
    clusters = {root: sorted(set(v)) for root, v in groups.items()}

    # annotate dataframe (nothing deleted)
    df["dup_cluster_id"] = ""
    df["dup_cluster_size"] = 0
    df["dup_role"] = ""
    df["dup_canonical_row"] = ""
    df["dup_match_embed_sim"] = np.nan
    df["dup_llm_reason"] = ""
    reason_by_node = {}
    for a, b, _s, r in confirmed:
        reason_by_node.setdefault(a, r); reason_by_node.setdefault(b, r)

    cross_release_clusters = 0
    cluster_rows = []
    for cid, (root, members) in enumerate(sorted(clusters.items()), start=1):
        def _score(i):
            ts = pd.to_numeric(df.loc[i].get("sightingDetails.trustScore"), errors="coerce")
            return (ts if pd.notna(ts) else 0, len(str(df.loc[i].get(mt._NARR, ""))))
        canon = max(members, key=_score)
        sources = sorted({df.loc[i, "_source_dataset"] for i in members})
        is_cross = len(sources) > 1
        cross_release_clusters += int(is_cross)
        for i in members:
            df.at[i, "dup_cluster_id"] = cid
            df.at[i, "dup_cluster_size"] = len(members)
            df.at[i, "dup_role"] = "canonical" if i == canon else "duplicate"
            df.at[i, "dup_canonical_row"] = canon
            df.at[i, "dup_match_embed_sim"] = round(sim_edge.get(i, np.nan), 4)
            df.at[i, "dup_llm_reason"] = reason_by_node.get(i, "")
        yrs = sorted({mt._int_or_none(df.loc[i].get("date_time.year")) for i in members})
        title = df.loc[canon].get("Title")
        if pd.isna(title) or not str(title).strip():
            title = df.loc[canon].get("document_id", "")
        cluster_rows.append({
            "cluster_id": cid, "size": len(members),
            "year(s)": ",".join(str(y) for y in yrs if y),
            "sources": ",".join(sources),
            "cross_release": is_cross,
            "canonical_row": canon,
            "canonical_title": str(title)[:60],
            "member_rows": ",".join(map(str, members)),
        })

    df["is_duplicate"] = df["dup_role"] == "duplicate"
    n_clusters = len(clusters)
    n_dups = int((df["dup_role"] == "duplicate").sum())
    print(f"clusters={n_clusters} (cross-release={cross_release_clusters}) | "
          f"rows in clusters={int((df['dup_cluster_id']!='').sum())} | "
          f"redundant (non-canonical) rows={n_dups}")

    df = _sanitize(df)
    csv_out = OUT_XLSX.replace(".xlsx", ".csv")
    df.to_csv(csv_out, index=False)
    print(f"wrote {csv_out} (raw, unsorted) {df.shape}")

    # ── re-read + shade + sort, mirroring color_dedup_xlsx.py exactly ──────
    flagged = _sanitize(pd.read_csv(csv_out, low_memory=False))
    flagged = flagged.sort_values("dup_cluster_id", ascending=False,
                                  na_position="last", kind="stable").reset_index(drop=True)
    clusters_df = pd.DataFrame(cluster_rows).sort_values("cluster_id", ascending=False)

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as xw:
        flagged.to_excel(xw, sheet_name="flagged", index=False)
        clusters_df.to_excel(xw, sheet_name="clusters", index=False)

    wb = openpyxl.load_workbook(OUT_XLSX)
    ws = wb["flagged"]
    col = {c.value: c.column for c in ws[1]}["dup_cluster_id"]
    ncol = ws.max_column
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    colored = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v in (None, ""):
            continue
        try:
            cid = int(v)
        except (TypeError, ValueError):
            continue
        fill = EVEN if cid % 2 == 0 else ODD
        for cc in range(1, ncol + 1):
            ws.cell(r, cc).fill = fill
        colored += 1
    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX} (sheets: flagged {flagged.shape}, clusters "
          f"[{n_clusters}]) | shaded {colored} rows")
    return n_clusters, n_dups, cross_release_clusters


if __name__ == "__main__":
    main()
